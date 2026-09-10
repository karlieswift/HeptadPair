#!/usr/bin/env python3
from pathlib import Path
import argparse
import json
import re
from openpyxl import load_workbook

ABS=[
    re.compile('/'+'home'+'/'),
    re.compile('/'+'mnt'+'/'),
    re.compile(r'\b[A-Za-z]:'+r'\\'),
    re.compile('hgq'+'@'),
    re.compile('as'+'-4125'),
    re.compile(r'r750-\d+'),
    re.compile('202001'+'Main'),
]

WORKBOOKS={
    'base': 'source_data/HeptadPair_Source_Data_v8.5_BASE_MANUSCRIPT_20260830.xlsx',
    'revision': 'source_data/HeptadPair_Source_Data_v8.9_GITHUB_20260903.xlsx',
}
ALLOWED_BASE_FORMULAS={
    'Feature_Accounting!C2',
    'Feature_Accounting!C3',
    'Feature_Accounting!C4',
}


def audit_workbook(path, allow_formulas=None):
    allow_formulas=set(allow_formulas or [])
    wb=load_workbook(path,data_only=False,read_only=True)
    formulas=[]; unexpected_formulas=[]; path_hits=[]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value,str):
                    continue
                coord=f'{ws.title}!{c.coordinate}'
                if c.value.startswith('='):
                    formulas.append(coord)
                    if coord not in allow_formulas:
                        unexpected_formulas.append(coord)
                if any(r.search(c.value) for r in ABS):
                    path_hits.append(coord)
    return {
        'path': str(path),
        'sheets': wb.sheetnames,
        'sheet_count': len(wb.sheetnames),
        'formula_cells': formulas,
        'unexpected_formula_cells': unexpected_formulas,
        'absolute_or_private_path_cells': path_hits,
        'pass': not unexpected_formulas and not path_hits,
    }


def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--repo',default='.'); args=ap.parse_args()
    root=Path(args.repo).resolve()
    out={}
    for key,rel in WORKBOOKS.items():
        p=root/rel
        if not p.is_file():
            out[key]={'path':str(p),'pass':False,'error':'missing workbook'}
            continue
        allow=ALLOWED_BASE_FORMULAS if key=='base' else set()
        out[key]=audit_workbook(p,allow)
    result={'workbooks':out,'pass':all(v.get('pass',False) for v in out.values())}
    print(json.dumps(result,indent=2))
    if not result['pass']:
        raise SystemExit(2)

if __name__=='__main__':
    main()
