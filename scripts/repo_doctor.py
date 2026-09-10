#!/usr/bin/env python3
from pathlib import Path
import argparse, json, re, zipfile
try:
    from openpyxl import load_workbook
except Exception:
    load_workbook=None

ABS_PATTERNS=[re.compile('/'+'home'+'/'),re.compile('/'+'mnt'+'/'),re.compile(r'\b[A-Za-z]:'+r'\\'),re.compile('hgq'+'@'),re.compile('as'+'-4125'),re.compile(r'r750-\d+'),re.compile('202001'+'Main')]
TEXT_EXTS={'.md','.txt','.json','.csv','.tsv','.py','.sh','.tex','.yml','.yaml','.cff','.pml','.bib','.bibtex','.fasta','.fa','.a3m','.pdb','.log'}
REQUIRED=['README.md','CITATION.cff','LICENSE_SELECTION_REQUIRED.md','VERSION','PUBLIC_RELEASE_GATE.json','MANIFEST.tsv','SHA256SUMS.txt',
          'manuscript/HeptadPair_Main_Manuscript_v8.9_PHASE81_PHASE82_20260903.pdf',
          'manuscript/HeptadPair_Main_Manuscript_v8.9_PHASE81_PHASE82_20260903.tex',
          'supplementary/HeptadPair_Supplementary_Information_v8.9_PHASE81_PHASE82_20260903.pdf',
          'source_data/HeptadPair_Source_Data_v8.5_BASE_MANUSCRIPT_20260830.xlsx',
          'source_data/HeptadPair_Source_Data_v8.9_GITHUB_20260903.xlsx']

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--repo',default='.'); ap.add_argument('--strict',action='store_true'); ap.add_argument('--require-public',action='store_true'); args=ap.parse_args()
    root=Path(args.repo).resolve(); issues=[]; warnings=[]
    for rel in REQUIRED:
        if not (root/rel).is_file(): issues.append(f'missing required file: {rel}')
    for p in root.rglob('*'):
        if '__pycache__' in p.parts or '.pytest_cache' in p.parts or p.suffix=='.pyc': issues.append(f'compiled/cache file present: {p.relative_to(root)}')
        if p.is_file() and p.stat().st_size>50*1024*1024: warnings.append(f'file exceeds 50 MiB: {p.relative_to(root)}')
        if p.is_file() and p.suffix.lower() in TEXT_EXTS:
            try: s=p.read_text(encoding='utf-8',errors='strict')
            except Exception: continue
            for pat in ABS_PATTERNS:
                if pat.search(s): issues.append(f'absolute/private path pattern in {p.relative_to(root)}: {pat.pattern}'); break
    if load_workbook:
        for xlsx_rel in [
            'source_data/HeptadPair_Source_Data_v8.5_BASE_MANUSCRIPT_20260830.xlsx',
            'source_data/HeptadPair_Source_Data_v8.9_GITHUB_20260903.xlsx',
        ]:
            xlsx=root/xlsx_rel
            if not xlsx.is_file():
                continue
            wb=load_workbook(xlsx,read_only=True,data_only=False)
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value,str):
                            for pat in ABS_PATTERNS:
                                if pat.search(c.value): issues.append(f'absolute/private path in workbook {xlsx_rel}:{ws.title}!{c.coordinate}'); break
    gate_path=root/'PUBLIC_RELEASE_GATE.json'
    reviewer=False; public=False
    if gate_path.is_file():
        gate=json.loads(gate_path.read_text()); reviewer=bool(gate.get('reviewer_ready')); public=bool(gate.get('public_release_ready'))
        if not reviewer: issues.append('PUBLIC_RELEASE_GATE reviewer_ready is false')
        if args.require_public and not public: issues.append('PUBLIC_RELEASE_GATE public_release_ready is false')
    result={'reviewer_ready':reviewer and not issues,'public_release_ready':public and not issues,'strict':args.strict,'issues':issues,'warnings':warnings}
    print(json.dumps(result,indent=2))
    if args.strict and issues: raise SystemExit(2)
if __name__=='__main__': main()
